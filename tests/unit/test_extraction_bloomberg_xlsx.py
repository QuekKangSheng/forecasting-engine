"""Reading Bloomberg .xlsx workbook exports into the open-schema shape.

Two layouts are built here: the one-sheet Spreadsheet Builder export, laid out
like the CSV export, and the older ``Data`` + ``Metadata`` workbook.
"""

import re
import zipfile
from datetime import datetime
from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from forecasting_engine.extraction import bloomberg_csv
from forecasting_engine.extraction.bloomberg_xlsx import (
    UNSAVED_MESSAGE,
    BloombergXlsxError,
    read_export,
)

DATES = ["2024-01-01", "2024-01-02", "2024-01-03"]


def workbook_bytes(
    *,
    security="SPX Index",
    fields=("PX_LAST",),
    rows=(("2024-01-01", 100.0),),
    data_sheet="Data",
    with_metadata=True,
) -> bytes:
    """A workbook shaped like the real Bloomberg exports, as raw bytes."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = data_sheet
    sheet.append(["Date", *fields])
    for row in rows:
        date, *values = row
        sheet.append([datetime.fromisoformat(date), *values])
    if with_metadata:
        meta = book.create_sheet("Metadata")
        meta.append(["Field", "Value"])
        meta.append(["Security", security])
        meta.append(["Period", "D"])
    return _saved(book)


def _saved(book: openpyxl.Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


ONE_SHEET_ROWS = [
    ["Security", "SPX Index", None],
    ["Start Date", datetime(2024, 1, 1), None],
    ["End Date", datetime(2024, 1, 4), None],
    ["Period", "D", None],
    ["Currency", "USD", None],
    [None, None, None],
    ["Date", "PX_LAST", "TOT_RETURN_INDEX_GROSS_DVDS"],
    [datetime(2024, 1, 2), 4742.83, 10870.5],
    [datetime(2024, 1, 3), 4704.81, "#N/A N/A"],
    [datetime(2024, 1, 4), 4688.68, 10751.25],
]

ONE_SHEET_CSV = b"""Security,SPX Index,
Start Date,1/1/2024,
End Date,1/4/2024,
Period,D,
Currency,USD,
,,
Date,PX_LAST,TOT_RETURN_INDEX_GROSS_DVDS
1/2/2024,4742.83,10870.5
1/3/2024,4704.81,#N/A N/A
1/4/2024,4688.68,10751.25
"""


def one_sheet_bytes(rows=ONE_SHEET_ROWS) -> bytes:
    """A Spreadsheet Builder export: one sheet, laid out like the CSV export."""
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    return _saved(book)


def with_saved_values(data: bytes, values: dict[str, object]) -> bytes:
    """``data`` with each formula cell named in ``values`` given that saved result,
    as Excel writes it after recalculating. openpyxl itself never saves one."""
    source = zipfile.ZipFile(BytesIO(data))
    sheet = "xl/worksheets/sheet1.xml"
    xml = source.read(sheet).decode()
    for ref, value in values.items():
        kind = ' t="str"' if isinstance(value, str) else ""
        xml, found = re.subn(
            rf'<c r="{ref}"([^>]*)><f>(.*?)</f><v ?/>(?:</v>)?</c>',
            lambda m, ref=ref, kind=kind, value=value: (
                f'<c r="{ref}"{m[1]}{kind}><f>{m[2]}</f><v>{value}</v></c>'
            ),
            xml,
        )
        assert found == 1, ref
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            target.writestr(item, xml if item.filename == sheet else source.read(item))
    return out.getvalue()


BDH = '=BDH("SPX Index","PX_LAST","1/2/2024","1/3/2024")'


def bdh_rows(date_cell=BDH):
    """The first sheet as Bloomberg's add-in leaves it: the table's top-left cell
    holds the BDH formula, and its results spill into the cells beside and below.
    A spilled date may be stored as an Excel serial number."""
    return [
        ["Security", "SPX Index"],
        ["Start Date", datetime(2024, 1, 2)],
        ["End Date", datetime(2024, 1, 3)],
        ["Period", "D"],
        [None, None],
        [date_cell, "PX_LAST"],
        ["=A6+0", 4742.83],
        [45294, 4704.81],
    ]


def test_security_comes_from_the_metadata_sheet_not_the_filename():
    data = workbook_bytes(security="VIX Index")

    export = read_export("misleading_spx_name.xlsx", data)

    assert export.security == "VIX Index"


def test_every_field_in_the_data_sheet_is_kept_not_just_one():
    data = workbook_bytes(
        security="SPX Index",
        fields=("PX_LAST", "PX_BID", "TOT_RETURN_INDEX_GROSS_DVDS"),
        rows=[("2024-01-01", 100.0, 99.9, 105.0)],
    )

    export = read_export("spx.xlsx", data)

    assert set(export.frame.columns) == {
        "Date",
        "SPX_Index_PX_LAST",
        "SPX_Index_PX_BID",
        "SPX_Index_TOT_RETURN_INDEX_GROSS_DVDS",
    }


def test_placeholder_values_become_nan_not_a_literal_string():
    data = workbook_bytes(
        security="JPMVXYGL Index",
        fields=("PX_LAST", "PX_BID"),
        rows=[("2024-01-01", 6.5, "#N/A N/A"), ("2024-01-02", 6.6, "#N/A N/A")],
    )

    export = read_export("jpm.xlsx", data)

    assert export.frame["JPMVXYGL_Index_PX_BID"].isna().all()
    assert export.frame["JPMVXYGL_Index_PX_LAST"].tolist() == [6.5, 6.6]


def test_a_one_sheet_export_reads_exactly_as_the_same_export_saved_as_csv():
    from_xlsx = read_export("spx.xlsx", one_sheet_bytes())
    from_csv = bloomberg_csv.read_export("spx.csv", ONE_SHEET_CSV)

    assert from_xlsx.security == from_csv.security == "SPX Index"
    pd.testing.assert_frame_equal(from_xlsx.frame, from_csv.frame)


def test_a_one_sheet_export_of_bdh_formulas_reads_their_saved_values():
    data = with_saved_values(one_sheet_bytes(bdh_rows()), {"A6": "Date", "A7": 45293})

    export = read_export("spx.xlsx", data)

    assert export.frame["Date"].tolist() == [pd.Timestamp("2024-01-02"), pd.Timestamp("2024-01-03")]
    assert export.frame["SPX_Index_PX_LAST"].tolist() == [4742.83, 4704.81]


def test_a_formula_with_no_saved_value_asks_for_the_file_to_be_saved_in_excel():
    data = with_saved_values(one_sheet_bytes(bdh_rows()), {"A6": "Date"})

    with pytest.raises(BloombergXlsxError, match=re.escape(UNSAVED_MESSAGE)):
        read_export("spx.xlsx", data)


def test_a_sheet_still_requesting_data_asks_for_the_file_to_be_saved_in_excel():
    data = one_sheet_bytes(bdh_rows(date_cell="#N/A Requesting Data..."))

    with pytest.raises(BloombergXlsxError, match=re.escape(UNSAVED_MESSAGE)):
        read_export("spx.xlsx", data)


def test_a_first_sheet_not_starting_with_security_names_the_expected_layout():
    data = workbook_bytes(data_sheet="Sheet1")

    with pytest.raises(BloombergXlsxError, match="Security / Start Date / End Date / Period"):
        read_export("odd.xlsx", data)


def test_a_workbook_without_metadata_still_reads_with_an_empty_security():
    data = workbook_bytes(with_metadata=False)

    export = read_export("mystery.xlsx", data)

    assert export.security == ""


def test_a_repeated_date_keeps_the_last_value_and_reports_a_note():
    data = workbook_bytes(
        security="A Index",
        fields=("PX_LAST",),
        rows=[("2024-01-01", 1.0), ("2024-01-01", 2.0), ("2024-01-02", 3.0)],
    )

    export = read_export("a.xlsx", data)

    assert export.frame["A_Index_PX_LAST"].tolist() == [2.0, 3.0]
    assert len(export.notes) == 1
    assert "repeated date" in export.notes[0]


def test_dates_are_parsed_as_real_datetimes():
    data = workbook_bytes(rows=[(d, 1.0) for d in DATES])

    export = read_export("spx.xlsx", data)

    assert pd.api.types.is_datetime64_any_dtype(export.frame["Date"])
    assert export.frame["Date"].tolist() == [pd.Timestamp(d) for d in DATES]


def test_not_a_real_xlsx_file_is_refused_not_crashed():
    with pytest.raises(BloombergXlsxError, match="not a readable .xlsx"):
        read_export("fake.xlsx", b"this is not a zip file")
